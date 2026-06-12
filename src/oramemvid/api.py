"""FastAPI REST API wiring all oramemvid modules together."""

import os
import tempfile
import zipfile
from contextlib import asynccontextmanager, contextmanager
from pathlib import PurePosixPath

from fastapi import FastAPI, HTTPException, Query, Request, UploadFile
from fastapi.responses import JSONResponse
from pydantic import BaseModel

from oramemvid.config import get_settings
from oramemvid.db import get_pool, close_pool, init_schema, get_capabilities
from oramemvid.embeddings import get_embedding_provider
from oramemvid.frames import get_frame, list_frames, delete_frame
from oramemvid.ingest import ingest_file, ingest_text
from oramemvid.llm import get_llm_provider
from oramemvid.memory_cards import (
    create_cards_from_llm,
    get_memory_card,
    list_memory_cards,
    delete_memory_card,
)
from oramemvid.search import SearchCapabilityError, search_text, search_vector, search_hybrid

settings = get_settings()
embedding_provider = get_embedding_provider(settings)
llm_provider = get_llm_provider(settings)
_UPLOAD_COPY_CHUNK_BYTES = 1024 * 1024
_MAX_OFFICE_ARCHIVE_MEMBERS = 1000
_MAX_OFFICE_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
_MAX_OFFICE_COMPRESSION_RATIO = 100
_MULTIPART_OVERHEAD_BYTES = 1024 * 1024


def _get_conn():
    return get_pool().acquire()


def _release_conn(conn):
    get_pool().release(conn)


@contextmanager
def _db_conn():
    conn = _get_conn()
    try:
        yield conn
    finally:
        _release_conn(conn)


def _validate_upload_suffix(filename: str | None) -> str:
    if not filename:
        raise HTTPException(status_code=415, detail="Uploaded file must have a filename")
    suffix = os.path.splitext(filename)[1].lower()
    if not suffix or suffix not in settings.upload_extension_set:
        allowed = ", ".join(sorted(settings.upload_extension_set))
        raise HTTPException(
            status_code=415,
            detail=f"Unsupported file type. Allowed extensions: {allowed}",
        )
    return suffix


def _write_upload_to_temp(file: UploadFile, suffix: str, max_bytes: int) -> str:
    total = 0
    tmp_path = None
    try:
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
            tmp_path = tmp.name
            while True:
                chunk = file.file.read(_UPLOAD_COPY_CHUNK_BYTES)
                if not chunk:
                    break
                total += len(chunk)
                if total > max_bytes:
                    raise HTTPException(
                        status_code=413,
                        detail=f"Uploaded file exceeds {max_bytes} bytes",
                    )
                tmp.write(chunk)
        return tmp_path
    except Exception:
        if tmp_path and os.path.exists(tmp_path):
            os.unlink(tmp_path)
        raise


def _validate_upload_content(file_path: str, suffix: str) -> None:
    try:
        with open(file_path, "rb") as f:
            header = f.read(8)
    except OSError as exc:
        raise HTTPException(status_code=400, detail="Could not read uploaded file") from exc

    if suffix == ".pdf":
        if not header.startswith(b"%PDF-"):
            raise HTTPException(status_code=415, detail="Uploaded file is not a PDF")
        _validate_pdf_parseable(file_path)
        return

    if suffix == ".txt":
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                for _ in f:
                    pass
        except UnicodeDecodeError as exc:
            raise HTTPException(status_code=415, detail="Uploaded text file must be UTF-8") from exc
        return

    office_expectations = {
        ".docx": ("word/", "word/document.xml"),
        ".xlsx": ("xl/", "xl/workbook.xml"),
        ".pptx": ("ppt/", "ppt/presentation.xml"),
    }
    expected = office_expectations.get(suffix)
    if expected is None:
        return
    expected_root, required_document = expected

    if not zipfile.is_zipfile(file_path):
        raise HTTPException(status_code=415, detail=f"Uploaded file is not a valid {suffix} archive")
    try:
        _validate_office_archive(file_path, suffix, expected_root, required_document)
    except zipfile.BadZipFile as exc:
        raise HTTPException(status_code=415, detail=f"Uploaded file is not a valid {suffix} archive") from exc
    _validate_office_parseable(file_path, suffix)


def _validate_office_archive(
    file_path: str,
    suffix: str,
    expected_root: str,
    required_document: str,
) -> None:
    total_uncompressed = 0
    saw_expected_root = False
    saw_content_types = False
    saw_relationships = False
    saw_required_document = False

    with zipfile.ZipFile(file_path) as archive:
        members = archive.infolist()
        if len(members) > _MAX_OFFICE_ARCHIVE_MEMBERS:
            raise HTTPException(status_code=413, detail="Office archive contains too many files")

        for member in members:
            name = member.filename
            parts = PurePosixPath(name).parts
            if name.startswith("/") or ".." in parts:
                raise HTTPException(status_code=415, detail="Office archive contains an unsafe path")

            saw_expected_root = saw_expected_root or name.startswith(expected_root)
            saw_content_types = saw_content_types or name == "[Content_Types].xml"
            saw_relationships = saw_relationships or name == "_rels/.rels"
            saw_required_document = saw_required_document or name == required_document
            total_uncompressed += member.file_size
            if total_uncompressed > _MAX_OFFICE_UNCOMPRESSED_BYTES:
                raise HTTPException(status_code=413, detail="Office archive expands too large")

            if (
                member.compress_size > 0
                and member.file_size / member.compress_size > _MAX_OFFICE_COMPRESSION_RATIO
            ):
                raise HTTPException(status_code=413, detail="Office archive compression ratio is too high")

    if not saw_expected_root:
        raise HTTPException(status_code=415, detail=f"Uploaded file does not match {suffix} content")
    if not (saw_content_types and saw_relationships and saw_required_document):
        raise HTTPException(status_code=415, detail=f"Uploaded file is missing required {suffix} structure")


def _validate_pdf_parseable(file_path: str) -> None:
    doc = None
    try:
        import pymupdf

        doc = pymupdf.open(file_path)
        if doc.needs_pass:
            raise ValueError("encrypted PDF is not supported")
    except Exception as exc:
        raise HTTPException(status_code=415, detail="Uploaded file is not a parseable PDF") from exc
    finally:
        if doc is not None:
            doc.close()


def _validate_office_parseable(file_path: str, suffix: str) -> None:
    try:
        if suffix == ".docx":
            from docx import Document

            Document(file_path)
        elif suffix == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(file_path, read_only=True)
            workbook.close()
        elif suffix == ".pptx":
            from pptx import Presentation

            Presentation(file_path)
    except Exception as exc:
        raise HTTPException(status_code=415, detail=f"Uploaded file is not a parseable {suffix} file") from exc


def _parse_tag_filters(tags: list[str] | None) -> dict[str, str] | None:
    if not tags:
        return None
    parsed: dict[str, str] = {}
    for raw in tags:
        for item in raw.split(","):
            item = item.strip()
            if not item:
                continue
            if "=" not in item:
                raise HTTPException(
                    status_code=400,
                    detail="Tag filters must use key=value syntax",
                )
            key, value = item.split("=", 1)
            key = key.strip()
            value = value.strip()
            if not key or not value:
                raise HTTPException(
                    status_code=400,
                    detail="Tag filters must include both key and value",
                )
            parsed[key] = value
    return parsed or None


def _search_filter_args(
    tags: list[str] | None,
    entity: str | None,
    kind: str | None,
) -> dict:
    return {
        "tags": _parse_tag_filters(tags),
        "entity": entity,
        "kind": kind,
    }


def _run_filtered_search(
    search_func,
    query: str,
    top_k: int,
    time_from: str | None,
    time_to: str | None,
    tags: list[str] | None,
    entity: str | None,
    kind: str | None,
    *extra_args,
):
    filter_args = _search_filter_args(tags, entity, kind)
    with _db_conn() as conn:
        try:
            return search_func(
                conn,
                query,
                *extra_args,
                top_k=top_k,
                time_from=time_from,
                time_to=time_to,
                **filter_args,
            )
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc
        except SearchCapabilityError as exc:
            raise HTTPException(status_code=503, detail=str(exc)) from exc


@asynccontextmanager
async def lifespan(app: FastAPI):
    pool = get_pool(settings)
    with pool.acquire() as conn:
        init_schema(conn, settings)
    yield
    close_pool()


app = FastAPI(title="oramemvid", version="0.1.0", lifespan=lifespan)


@app.middleware("http")
async def reject_large_upload_requests(request: Request, call_next):
    if request.url.path == "/ingest/file" and request.method == "POST":
        content_length = request.headers.get("content-length")
        limit = settings.max_upload_bytes + _MULTIPART_OVERHEAD_BYTES
        if content_length is None:
            return JSONResponse({"detail": "Content-Length header is required"}, status_code=411)
        try:
            request_size = int(content_length)
        except ValueError:
            return JSONResponse({"detail": "Invalid Content-Length header"}, status_code=400)
        if request_size > limit:
            return JSONResponse({"detail": f"Upload request exceeds {limit} bytes"}, status_code=413)
    return await call_next(request)


# --- Request models ---


class IngestTextRequest(BaseModel):
    text: str
    uri: str
    title: str | None = None
    extract_memories: bool = False


# --- Ingest routes ---


@app.post("/ingest/text")
def route_ingest_text(req: IngestTextRequest):
    with _db_conn() as conn:
        llm = llm_provider if req.extract_memories else None
        result = ingest_text(
            conn=conn,
            text=req.text,
            uri=req.uri,
            provider=embedding_provider,
            llm=llm,
            title=req.title,
            chunk_size=settings.chunk_size,
            chunk_overlap=settings.chunk_overlap,
        )
        return result


@app.post("/ingest/file")
def route_ingest_file(
    file: UploadFile,
    extract_memories: bool = Query(False),
):
    suffix = _validate_upload_suffix(file.filename)
    tmp_path = _write_upload_to_temp(file, suffix, settings.max_upload_bytes)
    try:
        _validate_upload_content(tmp_path, suffix)
        with _db_conn() as conn:
            llm = llm_provider if extract_memories else None
            result = ingest_file(
                conn=conn,
                file_path=tmp_path,
                provider=embedding_provider,
                llm=llm,
                chunk_size=settings.chunk_size,
                chunk_overlap=settings.chunk_overlap,
                filename=file.filename,
            )
            return result
    finally:
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)


# --- Frame routes ---


@app.get("/frames")
def route_list_frames(
    limit: int = Query(20, ge=1, le=1000),
    offset: int = Query(0, ge=0),
    doc_id: int | None = Query(None),
):
    with _db_conn() as conn:
        return list_frames(conn, limit=limit, offset=offset, doc_id=doc_id)


@app.get("/frames/{frame_id}")
def route_get_frame(frame_id: int):
    with _db_conn() as conn:
        frame = get_frame(conn, frame_id)
        if frame is None:
            raise HTTPException(status_code=404, detail="Frame not found")
        return frame


@app.delete("/frames/{frame_id}")
def route_delete_frame(frame_id: int):
    with _db_conn() as conn:
        deleted = delete_frame(conn, frame_id)
        if not deleted:
            raise HTTPException(status_code=404, detail="Frame not found")
        return {"deleted": True}


# --- Search routes ---


@app.get("/search")
def route_search(
    query: str = Query(...),
    mode: str = Query("hybrid", pattern="^(text|vector|hybrid)$"),
    top_k: int = Query(10, ge=1, le=100),
    time_from: str | None = Query(None),
    time_to: str | None = Query(None),
    tags: list[str] | None = Query(None, description="Tag filters as key=value"),
    entity: str | None = Query(None),
    kind: str | None = Query(None),
):
    if mode == "text":
        return _run_filtered_search(
            search_text, query, top_k, time_from, time_to, tags, entity, kind,
        )
    if mode == "vector":
        return _run_filtered_search(
            search_vector, query, top_k, time_from, time_to, tags, entity, kind,
            embedding_provider,
        )
    return _run_filtered_search(
        search_hybrid, query, top_k, time_from, time_to, tags, entity, kind,
        embedding_provider,
    )


@app.get("/search/text")
def route_search_text(
    query: str = Query(...),
    top_k: int = Query(10, ge=1, le=100),
    time_from: str | None = Query(None),
    time_to: str | None = Query(None),
    tags: list[str] | None = Query(None, description="Tag filters as key=value"),
    entity: str | None = Query(None),
    kind: str | None = Query(None),
):
    return _run_filtered_search(
        search_text, query, top_k, time_from, time_to, tags, entity, kind,
    )


@app.get("/search/vector")
def route_search_vector(
    query: str = Query(...),
    top_k: int = Query(10, ge=1, le=100),
    time_from: str | None = Query(None),
    time_to: str | None = Query(None),
    tags: list[str] | None = Query(None, description="Tag filters as key=value"),
    entity: str | None = Query(None),
    kind: str | None = Query(None),
):
    return _run_filtered_search(
        search_vector, query, top_k, time_from, time_to, tags, entity, kind,
        embedding_provider,
    )


# --- Memory card routes ---


@app.get("/memory")
def route_list_memory_cards(
    entity: str | None = Query(None),
    kind: str | None = Query(None),
    source_frame_id: int | None = Query(None),
    limit: int = Query(50, ge=1, le=1000),
    offset: int = Query(0, ge=0),
):
    with _db_conn() as conn:
        return list_memory_cards(
            conn, entity=entity, kind=kind,
            source_frame_id=source_frame_id, limit=limit, offset=offset,
        )


@app.get("/memory/{card_id}")
def route_get_memory_card(card_id: int):
    with _db_conn() as conn:
        card = get_memory_card(conn, card_id)
        if card is None:
            raise HTTPException(status_code=404, detail="Memory card not found")
        return card


@app.post("/memory/extract/{frame_id}")
def route_extract_memories(frame_id: int):
    with _db_conn() as conn:
        frame = get_frame(conn, frame_id)
        if frame is None:
            raise HTTPException(status_code=404, detail="Frame not found")

        card_ids = create_cards_from_llm(
            conn, llm_provider, frame["content"], frame_id,
        )
        return {"frame_id": frame_id, "cards_created": len(card_ids), "card_ids": card_ids}


@app.delete("/memory/{card_id}")
def route_delete_memory_card(card_id: int):
    with _db_conn() as conn:
        deleted = delete_memory_card(conn, card_id)
        if not deleted:
            raise HTTPException(status_code=404, detail="Memory card not found")
        return {"deleted": True}


# --- Document routes ---


@app.get("/documents")
def route_list_documents(
    limit: int = Query(20, ge=1, le=1000),
    offset: int = Query(0, ge=0),
):
    with _db_conn() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            SELECT doc_id, filename, doc_type, file_hash, total_frames, ingested_at
            FROM documents
            ORDER BY doc_id
            OFFSET :off ROWS FETCH NEXT :lim ROWS ONLY
        """, {"off": offset, "lim": limit})
        return [
            {
                "doc_id": r[0], "filename": r[1], "doc_type": r[2],
                "file_hash": r[3], "total_frames": r[4],
                "ingested_at": r[5].isoformat() if r[5] else None,
            }
            for r in cursor.fetchall()
        ]


# --- Stats ---


@app.get("/stats")
def route_stats():
    with _db_conn() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT COUNT(*) FROM frames")
        frame_count = cursor.fetchone()[0]
        cursor.execute("SELECT COUNT(*) FROM memory_cards")
        card_count = cursor.fetchone()[0]
        cursor.execute("SELECT COUNT(*) FROM documents")
        document_count = cursor.fetchone()[0]
        return {
            "frame_count": frame_count,
            "card_count": card_count,
            "document_count": document_count,
        }


# --- Health ---


@app.get("/health")
def route_health():
    try:
        with _db_conn() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT 1 FROM dual")
            cursor.fetchone()
            capabilities = get_capabilities()
            return {
                "status": "degraded" if capabilities.degraded else "ok",
                "database": "connected",
                "capabilities": capabilities.as_dict(),
            }
    except Exception as exc:
        return {"status": "degraded", "database": str(exc)}
